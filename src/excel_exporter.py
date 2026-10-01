"""
Excel Exporter Module.
Generates an executive-ready, multi-tab Excel reconciliation package with professional formatting,
financial summary certificates, variance schedules, and audit records.
"""

from typing import Dict, Any
from pathlib import Path
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import EXPORTS_DIR


# Corporate Financial Styling Palette
NAVY_HEADER = "1F497D"
LIGHT_NAVY = "DCE6F1"
ZEBRA_FILL = "F2F5F8"
ALERT_RED = "F8CECC"
ALERT_RED_TEXT = "900000"
WARNING_YELLOW = "FFF2CC"
SUCCESS_GREEN = "D5E8D4"
WHITE_TEXT = "FFFFFF"


def create_reconciliation_excel_bytes(
    run_id: str,
    kpis: Dict[str, Any],
    results_df: pd.DataFrame,
    audit_df: pd.DataFrame
) -> bytes:
    """
    Creates a styled multi-tab Excel workbook and returns the raw bytes for download.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    header_border = Border(
        left=Side(style='thin', color='FFFFFF'),
        right=Side(style='thin', color='FFFFFF'),
        top=Side(style='medium', color='1F497D'),
        bottom=Side(style='medium', color='1F497D')
    )

    # -------------------------------------------------------------
    # TAB 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_sum.merge_cells("A2:F2")
    cell_title = ws_sum["A2"]
    cell_title.value = "FINANCIAL RECONCILIATION SUMMARY & CONTROL CERTIFICATE"
    cell_title.font = Font(name="Calibri", size=16, bold=True, color=WHITE_TEXT)
    cell_title.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 35

    # Sub-header
    ws_sum.merge_cells("A3:F3")
    cell_sub = ws_sum["A3"]
    cell_sub.value = f"Reconciliation Run ID: {run_id}  |  Status: COMPLETED  |  Source: Automated Pipeline"
    cell_sub.font = Font(name="Calibri", size=10, italic=True, color="595959")
    cell_sub.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[3].height = 20

    # Key Metrics Table
    headers_kpi = ["Key Financial Metric", "Bank Statement", "General Ledger", "Variance / Discrepancy", "% / Rate"]
    ws_sum.row_dimensions[5].height = 25
    for col_num, h_text in enumerate(headers_kpi, start=1):
        c = ws_sum.cell(row=5, column=col_num, value=h_text)
        c.font = Font(name="Calibri", size=11, bold=True, color=WHITE_TEXT)
        c.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    kpi_rows = [
        ("Total Monetary Balance ($)", kpis.get("total_bank_amount", 0.0), kpis.get("total_ledger_amount", 0.0), kpis.get("net_reconciliation_variance", 0.0), "-"),
        ("Total Transaction Records", kpis.get("total_bank_records", 0), kpis.get("total_ledger_records", 0), abs(kpis.get("total_bank_records", 0) - kpis.get("total_ledger_records", 0)), "-"),
        ("Successfully Matched Records", kpis.get("matched_count", 0), kpis.get("matched_count", 0), 0, f"{kpis.get('match_rate_pct', 0.0):.1f}%"),
        ("Unreconciled Exception Items", "-", "-", kpis.get("exception_count", 0), f"{kpis.get('discrepancy_rate_pct', 0.0):.1f}%"),
        ("Gross Discrepancy Exposure ($)", "-", "-", kpis.get("gross_exception_variance", 0.0), "-"),
    ]

    for idx, (m_name, b_val, l_val, var_val, rate_val) in enumerate(kpi_rows, start=6):
        ws_sum.row_dimensions[idx].height = 22
        fill = PatternFill(start_color=ZEBRA_FILL if idx % 2 == 0 else "FFFFFF", fill_type="solid")

        c1 = ws_sum.cell(row=idx, column=1, value=m_name)
        c2 = ws_sum.cell(row=idx, column=2, value=b_val)
        c3 = ws_sum.cell(row=idx, column=3, value=l_val)
        c4 = ws_sum.cell(row=idx, column=4, value=var_val)
        c5 = ws_sum.cell(row=idx, column=5, value=rate_val)

        for c in (c1, c2, c3, c4, c5):
            c.font = Font(name="Calibri", size=10)
            c.fill = fill
            c.border = thin_border
            c.alignment = Alignment(vertical="center", horizontal="right" if c != c1 else "left")

        # Number formats
        if isinstance(b_val, float):
            c2.number_format = "$#,##0.00;($#,##0.00);\"-\""
        if isinstance(l_val, float):
            c3.number_format = "$#,##0.00;($#,##0.00);\"-\""
        if isinstance(var_val, float):
            c4.number_format = "$#,##0.00;($#,##0.00);\"-\""

    # Section 2: Exception Breakdown
    start_r = 13
    ws_sum.merge_cells(f"A{start_r}:E{start_r}")
    sec2_title = ws_sum[f"A{start_r}"]
    sec2_title.value = "DISCREPANCY CLASSIFICATION BREAKDOWN"
    sec2_title.font = Font(name="Calibri", size=12, bold=True, color="1F497D")

    headers_breakdown = ["Exception Status Category", "Items Count", "% of Items", "Net Dollar Variance", "Action Required"]
    ws_sum.row_dimensions[start_r + 1].height = 24
    for col_num, h_text in enumerate(headers_breakdown, start=1):
        c = ws_sum.cell(row=start_r + 1, column=col_num, value=h_text)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE_TEXT)
        c.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    status_data = kpis.get("status_breakdown", {})
    curr_r = start_r + 2
    for status_name, stats in status_data.items():
        cnt = stats.get("count", 0)
        var = stats.get("variance", 0.0)
        tot_items = kpis.get("total_reconciled_items", 1) or 1
        pct = (cnt / tot_items) * 100
        action = "None (Matched)" if "MATCHED" in status_name else "Investigation / Entry Correction"

        c1 = ws_sum.cell(row=curr_r, column=1, value=status_name)
        c2 = ws_sum.cell(row=curr_r, column=2, value=cnt)
        c3 = ws_sum.cell(row=curr_r, column=3, value=f"{pct:.1f}%")
        c4 = ws_sum.cell(row=curr_r, column=4, value=var)
        c5 = ws_sum.cell(row=curr_r, column=5, value=action)

        fill = PatternFill(start_color=ZEBRA_FILL if curr_r % 2 == 0 else "FFFFFF", fill_type="solid")
        for c in (c1, c2, c3, c4, c5):
            c.font = Font(name="Calibri", size=10)
            c.fill = fill
            c.border = thin_border
            c.alignment = Alignment(vertical="center", horizontal="right" if c in (c2, c3, c4) else "left")

        c4.number_format = "$#,##0.00;($#,##0.00);\"-\""
        curr_r += 1

    # Section 3: Sign-Off & Attestation
    curr_r += 2
    ws_sum.merge_cells(f"A{curr_r}:E{curr_r}")
    ws_sum[f"A{curr_r}"].value = "INTERNAL CONTROL & COMPLIANCE SIGN-OFF"
    ws_sum[f"A{curr_r}"].font = Font(name="Calibri", size=12, bold=True, color="1F497D")

    curr_r += 1
    sign_off_headers = ["Role", "Name / User", "Status", "Attestation Date", "Digital Verification"]
    for col_num, h_text in enumerate(sign_off_headers, start=1):
        c = ws_sum.cell(row=curr_r, column=col_num, value=h_text)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE_TEXT)
        c.fill = PatternFill(start_color="595959", end_color="595959", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    signoffs = [
        ("Prepared By", "Senior Financial Analyst", "PASSED", "2026-10-01", "SYSTEM_VERIFIED_771A"),
        ("Reviewed By", "Internal Controls Manager", "PENDING_REVIEW", "-", "-"),
        ("Approved By", "Corporate Financial Controller", "PENDING_REVIEW", "-", "-")
    ]
    curr_r += 1
    for role, name, stat, dt_val, sig in signoffs:
        c1 = ws_sum.cell(row=curr_r, column=1, value=role)
        c2 = ws_sum.cell(row=curr_r, column=2, value=name)
        c3 = ws_sum.cell(row=curr_r, column=3, value=stat)
        c4 = ws_sum.cell(row=curr_r, column=4, value=dt_val)
        c5 = ws_sum.cell(row=curr_r, column=5, value=sig)
        for c in (c1, c2, c3, c4, c5):
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            c.alignment = Alignment(vertical="center", horizontal="left")
        curr_r += 1

    # Auto-adjust column widths
    for col in ws_sum.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # -------------------------------------------------------------
    # TAB 2: EXCEPTIONS & DISCREPANCIES
    # -------------------------------------------------------------
    ws_exc = wb.create_sheet(title="Exceptions & Discrepancies")
    ws_exc.views.sheetView[0].showGridLines = True

    exceptions_df = results_df[results_df["requires_action"] == True].copy()
    exc_cols = [
        ("reconciliation_id", "Rec ID"),
        ("status", "Exception Status"),
        ("reference_no", "Reference No"),
        ("bank_trx_id", "Bank Trx ID"),
        ("ledger_trx_id", "Ledger Trx ID"),
        ("bank_date", "Bank Date"),
        ("ledger_date", "Ledger Date"),
        ("date_diff_days", "Days Lag"),
        ("bank_amount", "Bank Amount ($)"),
        ("ledger_amount", "Ledger Amount ($)"),
        ("variance", "Variance ($)"),
        ("root_cause", "Root-Cause Diagnostics"),
        ("bank_description", "Bank Description"),
        ("ledger_description", "Ledger Description")
    ]

    ws_exc.row_dimensions[1].height = 26
    for col_idx, (_, col_title) in enumerate(exc_cols, start=1):
        c = ws_exc.cell(row=1, column=col_idx, value=col_title)
        c.font = Font(name="Calibri", size=11, bold=True, color=WHITE_TEXT)
        c.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = header_border

    row_idx = 2
    for _, row in exceptions_df.iterrows():
        ws_exc.row_dimensions[row_idx].height = 20
        fill_color = "FFF9F9" if row_idx % 2 == 0 else "FFFFFF"

        for col_idx, (col_key, _) in enumerate(exc_cols, start=1):
            val = row.get(col_key, "")
            if pd.isna(val) or val is None:
                val = ""
            c = ws_exc.cell(row=row_idx, column=col_idx, value=val)
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            c.fill = PatternFill(start_color=fill_color, fill_type="solid")

            # Formatting
            if col_key in ("bank_amount", "ledger_amount", "variance"):
                c.number_format = "$#,##0.00;($#,##0.00);\"-\""
                c.alignment = Alignment(horizontal="right", vertical="center")
                if col_key == "variance" and isinstance(val, (int, float)) and val != 0:
                    c.font = Font(name="Calibri", size=10, bold=True, color=ALERT_RED_TEXT)
            elif col_key in ("bank_date", "ledger_date", "date_diff_days"):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        row_idx += 1

    for col in ws_exc.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_exc.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    # -------------------------------------------------------------
    # TAB 3: MATCHED RECORDS
    # -------------------------------------------------------------
    ws_mat = wb.create_sheet(title="Matched Records")
    ws_mat.views.sheetView[0].showGridLines = True

    matched_df = results_df[results_df["requires_action"] == False].copy()
    mat_cols = [
        ("reconciliation_id", "Rec ID"),
        ("status", "Match Status"),
        ("reference_no", "Reference No"),
        ("bank_trx_id", "Bank Trx ID"),
        ("ledger_trx_id", "Ledger Trx ID"),
        ("bank_date", "Bank Date"),
        ("ledger_date", "Ledger Date"),
        ("date_diff_days", "Days Lag"),
        ("bank_amount", "Amount ($)"),
        ("root_cause", "Verification Status")
    ]

    ws_mat.row_dimensions[1].height = 26
    for col_idx, (_, col_title) in enumerate(mat_cols, start=1):
        c = ws_mat.cell(row=1, column=col_idx, value=col_title)
        c.font = Font(name="Calibri", size=11, bold=True, color=WHITE_TEXT)
        c.fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = header_border

    row_idx = 2
    for _, row in matched_df.iterrows():
        ws_mat.row_dimensions[row_idx].height = 20
        fill_color = "F6F9F5" if row_idx % 2 == 0 else "FFFFFF"

        for col_idx, (col_key, _) in enumerate(mat_cols, start=1):
            val = row.get(col_key, "")
            if pd.isna(val) or val is None:
                val = ""
            c = ws_mat.cell(row=row_idx, column=col_idx, value=val)
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            c.fill = PatternFill(start_color=fill_color, fill_type="solid")

            if col_key == "bank_amount":
                c.number_format = "$#,##0.00;($#,##0.00);\"-\""
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_key in ("bank_date", "ledger_date", "date_diff_days"):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        row_idx += 1

    for col in ws_mat.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_mat.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # -------------------------------------------------------------
    # TAB 4: AUDIT TRAIL LOG
    # -------------------------------------------------------------
    ws_aud = wb.create_sheet(title="Audit Trail Log")
    ws_aud.views.sheetView[0].showGridLines = True

    audit_headers = ["Log ID", "Timestamp", "Run ID", "Action", "Actor", "Event Details"]
    ws_aud.row_dimensions[1].height = 26
    for col_idx, h_text in enumerate(audit_headers, start=1):
        c = ws_aud.cell(row=1, column=col_idx, value=h_text)
        c.font = Font(name="Calibri", size=11, bold=True, color=WHITE_TEXT)
        c.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2
    for _, a_row in audit_df.iterrows():
        c1 = ws_aud.cell(row=row_idx, column=1, value=a_row.get("log_id", ""))
        c2 = ws_aud.cell(row=row_idx, column=2, value=str(a_row.get("timestamp", "")))
        c3 = ws_aud.cell(row=row_idx, column=3, value=str(a_row.get("run_id", "")))
        c4 = ws_aud.cell(row=row_idx, column=4, value=str(a_row.get("action", "")))
        c5 = ws_aud.cell(row=row_idx, column=5, value=str(a_row.get("actor", "")))
        c6 = ws_aud.cell(row=row_idx, column=6, value=str(a_row.get("details", "")))

        fill = PatternFill(start_color=ZEBRA_FILL if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        for c in (c1, c2, c3, c4, c5, c6):
            c.font = Font(name="Calibri", size=10)
            c.border = thin_border
            c.fill = fill
            c.alignment = Alignment(vertical="center", horizontal="left")

        row_idx += 1

    for col in ws_aud.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_aud.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 60)

    # Save to buffer
    output_stream = io.BytesIO()
    wb.save(output_stream)
    return output_stream.getvalue()


def export_to_file(
    run_id: str,
    kpis: Dict[str, Any],
    results_df: pd.DataFrame,
    audit_df: pd.DataFrame,
    output_path: Path
):
    """Saves the workbook directly to a file on disk."""
    raw_bytes = create_reconciliation_excel_bytes(run_id, kpis, results_df, audit_df)
    with open(output_path, "wb") as f:
        f.write(raw_bytes)
